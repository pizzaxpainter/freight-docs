"""Engine tests using SYNTHETIC files (the real PDFHandler.zip sample is not available here).

Run:  python -m unittest discover -s tests -v
Needs LibreOffice (soffice), poppler (pdftotext/pdftoppm) and qpdf on the PATH.
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
import openpyxl
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

import settings as S
from classifier import classify
from core import Workspace, extract_page_texts, validate_hawb, zip_outputs

TMP = Path(tempfile.mkdtemp(prefix="fdp_test_"))


def make_pdf(pages: list[list[str]]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            c.drawString(60, y, line)
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


def invoice(no="123", extra=""):
    return make_pdf([["COMMERCIAL INVOICE", f"Invoice No. {no}", "Unit Price   Qty   Amount", f"Total 1,234.56 {extra}"]])


def packing(no="123"):
    return make_pdf([["PACKING LIST", f"Invoice No. {no}", "Carton No   Gross Weight   Net Weight", "No. of cartons 10"]])


def make_zip(files: dict[str, bytes], encrypted_flag=False) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
        if encrypted_flag:
            for i in zf.infolist():
                i.flag_bits |= 0x1
    return buf.getvalue()


def soffice_convert(src: Path, fmt: str) -> Path:
    subprocess.run(["soffice", f"-env:UserInstallation={(TMP / 'prof').as_uri()}", "--headless", "--convert-to", fmt,
                    "--outdir", str(TMP), str(src)], check=True, capture_output=True, timeout=180)
    return TMP / f"{src.stem}.{fmt}"


def build_office_fixtures():
    d = docx.Document()
    d.add_heading("Packing List", 0)
    d.add_paragraph("Carton No 1   Gross Weight 12 kg   Net Weight 10 kg")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Qty"
    t.cell(1, 0).text, t.cell(1, 1).text = "Widget", "42"
    p = TMP / "pl_word.docx"
    d.save(p)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Billing"
    ws["A1"], ws["A2"], ws["A3"] = "Billing Report 4054522303", "Total charges", 98765.43
    wb.create_sheet("Second")["A1"] = "Second sheet content"
    p2 = TMP / "billing_report.xlsx"
    wb.save(p2)
    legacy_xls = soffice_convert(p2, "xls")
    legacy_doc = soffice_convert(p, "doc")
    return {"docx": p.read_bytes(), "xlsx": p2.read_bytes(), "xls": legacy_xls.read_bytes(), "doc": legacy_doc.read_bytes()}


def run_all(ws: Workspace, name: str, data: bytes):
    new, msgs = ws.ingest(name, data)
    for d in new:
        ws.process(d)
    return new, msgs


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = TMP / "sessions"
        if not hasattr(Base, "office"):
            Base.office = build_office_fixtures()

    def setUp(self):
        self.ws = Workspace(self.root)

    def tearDown(self):
        self.ws.cleanup()

    def text_of(self, path):
        return "\n".join(extract_page_texts(Path(path)))


class TestClassifier(unittest.TestCase):
    def test_content_beats_misleading_filename(self):
        r = classify("PACKING LIST\nCarton No 1\nGross Weight 5", "Invoice 999.pdf")
        self.assertEqual(r.category, "PL")
        self.assertEqual(r.confidence, "high")

    def test_generic_filename_content_decides(self):
        r = classify("COMMERCIAL INVOICE\nInvoice No. 55\nUnit Price", "20465890.pdf")
        self.assertEqual((r.category, r.needs_review), ("CI", False))

    def test_billing_report_requires_confirmation(self):
        r = classify("Billing Report 4054522303\nTotal charges", "Billing Report 4054522303.xls")
        self.assertEqual(r.category, "ID")
        self.assertTrue(r.needs_review)

    def test_unknown_is_needs_review_not_a_guess(self):
        r = classify("Hello world, nothing relevant here.", "scan0001.pdf")
        self.assertIsNone(r.category)
        self.assertTrue(r.needs_review)

    def test_filename_only_is_low_confidence_review(self):
        r = classify("some opaque text", "Packing List A.pdf")
        self.assertEqual((r.category, r.confidence, r.needs_review), ("PL", "low", True))

    def test_hawb_validation(self):
        self.assertEqual(validate_hawb(" sin01234567 ")[0], "SIN01234567")
        self.assertTrue(validate_hawb("")[1])
        self.assertTrue(validate_hawb("../evil")[1])
        self.assertTrue(validate_hawb("AB1")[1])
        self.assertTrue(validate_hawb("ABCD123")[2])   # allowed, but warned


class TestIngest(Base):
    def test_nested_zip_junk_and_unsafe_paths(self):
        z = make_zip({
            "Shipment 1/Invoice A.pdf": invoice("A"),
            "Shipment 1/sub/Packing A.pdf": packing("A"),
            "__MACOSX/Shipment 1/._Invoice A.pdf": b"junk",
            ".DS_Store": b"x",
            "Shipment 1/~$temp.xlsx": b"x",
            "../evil.pdf": invoice("evil"),
            "/abs/evil.pdf": invoice("evil2"),
            "Shipment 1/readme.txt": b"hello",
        })
        new, msgs = run_all(self.ws, "cust.zip", z)
        self.assertEqual(sorted(d.original_name for d in new), ["Invoice A.pdf", "Packing A.pdf"])
        text = " ".join(m[1] for m in msgs)
        self.assertIn("unsafe file path", text)
        self.assertIn("readme.txt", text)   # unsupported type is reported, not silently dropped
        self.assertTrue(all("cust.zip/Shipment 1" in d.source for d in new))
        for d in new:                       # nothing escaped the workspace
            self.assertTrue(str(d.orig_path).startswith(str(self.ws.dir)))
        self.assertFalse((self.root / "evil.pdf").exists())

    def test_duplicate_content_skipped_same_name_kept(self):
        a, b = invoice("1"), invoice("2")
        new1, _ = run_all(self.ws, "Invoice.pdf", a)
        new2, msgs = run_all(self.ws, "again.pdf", a)         # identical bytes
        new3, _ = run_all(self.ws, "Invoice.pdf", b)           # same filename, different content
        self.assertEqual((len(new1), len(new2), len(new3)), (1, 0, 1))
        self.assertIn("duplicate", msgs[0][1].lower())
        self.assertEqual(len(self.ws.docs), 2)

    def test_bad_and_protected_zips(self):
        _, m1 = run_all(self.ws, "bad.zip", b"this is not a zip")
        _, m2 = run_all(self.ws, "locked.zip", make_zip({"a.pdf": invoice()}, encrypted_flag=True))
        _, m3 = run_all(self.ws, "empty.zip", make_zip({}))
        _, m4 = run_all(self.ws, "notes.txt", b"hi")
        self.assertIn("not a valid ZIP", m1[0][1])
        self.assertIn("password-protected", m2[0][1])
        self.assertIn("no supported documents", m3[0][1])
        self.assertIn("unsupported file type", m4[0][1])
        self.assertEqual(self.ws.docs, [])

    def test_extraction_size_limit(self):
        old = S.MAX_EXTRACTED_BYTES
        S.MAX_EXTRACTED_BYTES = 1000
        try:
            new, msgs = run_all(self.ws, "big.zip", make_zip({"a.pdf": invoice()}))
        finally:
            S.MAX_EXTRACTED_BYTES = old
        self.assertEqual(new, [])
        self.assertEqual(msgs[0][0], "error")

    def test_corrupted_pdf_is_flagged_and_blocks_generation(self):
        new, _ = run_all(self.ws, "broken.pdf", b"%PDF-1.4\n garbage garbage")
        self.assertEqual(new[0].status, "Conversion Failed")
        good, _ = run_all(self.ws, "Invoice.pdf", invoice())
        out, errs = self.ws.generate("SIN01234567")
        self.assertEqual(out, [])
        self.assertTrue(any("broken.pdf" in e for e in errs))
        self.ws.remove(new[0].id)
        out, errs = self.ws.generate("SIN01234567")
        self.assertEqual(errs, [])
        self.assertEqual([o.filename for o in out], ["SIN01234567_CI.pdf"])

    def test_not_a_pdf_with_pdf_extension(self):
        new, _ = run_all(self.ws, "fake.pdf", b"PK\x03\x04not really")
        self.assertEqual(new[0].status, "Conversion Failed")


class TestConversion(Base):
    def test_all_office_formats_convert_with_content_preserved(self):
        d_docx, _ = run_all(self.ws, "Packing List word.docx", self.office["docx"])
        d_doc, _ = run_all(self.ws, "old packing.doc", self.office["doc"])
        d_xlsx, _ = run_all(self.ws, "billing.xlsx", self.office["xlsx"])
        d_xls, _ = run_all(self.ws, "Billing Report 4054522303.xls", self.office["xls"])
        for d in (d_docx[0], d_doc[0], d_xlsx[0], d_xls[0]):
            self.assertIsNone(d.error, d.error)
            self.assertIsNotNone(d.pdf_path)
            self.assertGreaterEqual(d.pages, 1)
        self.assertIn("Widget", self.text_of(d_docx[0].pdf_path))
        self.assertIn("Widget", self.text_of(d_doc[0].pdf_path))
        t = self.text_of(d_xls[0].pdf_path)
        self.assertIn("98765.43", t.replace(",", ""))
        self.assertIn("Second sheet content", t)           # second worksheet not silently dropped
        self.assertEqual(d_docx[0].category, "PL")
        self.assertEqual(d_xls[0].category, "ID")
        self.assertTrue(d_xls[0].needs_review)             # billing report must be confirmed
        self.assertTrue(any("worksheets" in w for w in d_xlsx[0].warnings))

    def test_garbage_office_file_fails_cleanly(self):
        new, _ = run_all(self.ws, "bad.xlsx", b"PK\x03\x04 this is not an xlsx at all")
        self.assertEqual(new[0].status, "Conversion Failed")
        self.assertTrue(new[0].error)


class TestConsolidation(Base):
    def test_multi_zip_plus_loose_files_merge_per_category_in_order(self):
        z1 = make_zip({"Invoice A.pdf": invoice("A"), "Packing List A.pdf": packing("A"), "Internal Document A.pdf":
                       make_pdf([["INTERNAL DOCUMENT", "Ref A"]])})
        z2 = make_zip({"f/Invoice B.pdf": invoice("B"), "f/Packing B.docx": self.office["docx"]})
        z3 = make_zip({"Invoice C.pdf": invoice("C"), "Packing List C.pdf": packing("C"),
                       "Internal Document C.pdf": make_pdf([["INTERNAL DOCUMENT", "Ref C"]])})
        for n, z in (("one.zip", z1), ("two.zip", z2), ("three.zip", z3)):
            run_all(self.ws, n, z)
        run_all(self.ws, "Invoice D.pdf", invoice("D"))
        self.assertEqual(len(self.ws.docs), 9)
        for d in self.ws.docs:
            if d.needs_review:
                self.ws.confirm(d.id)
        self.assertEqual(self.ws.blocking_issues(), [])
        outs, errs = self.ws.generate("sin01234567")
        self.assertEqual(errs, [])
        names = [o.filename for o in outs]
        self.assertEqual(names, ["SIN01234567_CI.pdf", "SIN01234567_PL.pdf", "SIN01234567_ID.pdf"])
        by = {o.category: o for o in outs}
        self.assertEqual((by["CI"].pages, by["PL"].pages, by["ID"].pages), (4, 3, 2))
        ci_text = self.text_of(by["CI"].path)
        order = [ci_text.index(f"Invoice No. {x}") for x in "ABCD"]
        self.assertEqual(order, sorted(order))                 # predictable order = upload order
        for o in outs:                                         # valid PDFs
            self.assertEqual(len(PdfReader(str(o.path)).pages), o.pages)
            self.assertEqual(subprocess.run(["qpdf", "--check", str(o.path)], capture_output=True).returncode, 0)
        self.assertNotIn("SIN01234567_DG.pdf", names)          # no empty files for missing categories
        z = zipfile.ZipFile(io.BytesIO(zip_outputs(outs)))
        self.assertEqual(sorted(z.namelist()), sorted(names))

    def test_reorder_and_stale_outputs(self):
        run_all(self.ws, "Invoice A.pdf", invoice("A"))
        run_all(self.ws, "Invoice B.pdf", invoice("B"))
        a, b = self.ws.docs
        self.ws.move(b.id, -1)
        outs, _ = self.ws.generate("SIN01234567")
        t = self.text_of(outs[0].path)
        self.assertLess(t.index("Invoice No. B"), t.index("Invoice No. A"))
        self.ws.remove(a.id)                                   # changing documents invalidates outputs
        self.assertEqual(self.ws.outputs, [])

    def test_generation_blocked_until_reviewed_and_hawb_valid(self):
        run_all(self.ws, "scan0001.pdf", make_pdf([["unrelated text"]]))
        d = self.ws.docs[0]
        self.assertEqual(d.status, "Needs Review")
        _, errs = self.ws.generate("SIN01234567")
        self.assertTrue(any("no category" in e for e in errs))
        self.ws.set_category(d.id, "PL")
        _, errs = self.ws.generate("")
        self.assertTrue(any("HAWB" in e for e in errs))
        _, errs = self.ws.generate("../../etc")
        self.assertTrue(errs)
        outs, errs = self.ws.generate("SIN01234567")
        self.assertEqual((errs, [o.filename for o in outs]), ([], ["SIN01234567_PL.pdf"]))

    def test_empty_session_cannot_generate(self):
        _, errs = self.ws.generate("SIN01234567")
        self.assertTrue(errs)

    def test_large_batch(self):
        for i in range(60):
            run_all(self.ws, f"Invoice {i}.pdf", invoice(str(i)))
        outs, errs = self.ws.generate("SIN01234567")
        self.assertEqual((errs, outs[0].pages), ([], 60))


class TestSplit(Base):
    def mixed(self):
        return make_pdf([["COMMERCIAL INVOICE", "Invoice No. 7", "Unit Price"],
                         ["continuation of the invoice, line items only"],
                         ["PACKING LIST", "Carton No 1", "Gross Weight 3"],
                         ["more packing detail"]])

    def test_mixed_detected_split_and_no_pages_lost(self):
        new, _ = run_all(self.ws, "customer.pdf", self.mixed())
        d = new[0]
        self.assertTrue(d.mixed)
        self.assertEqual(d.status, "Needs Review")
        self.assertEqual(self.ws.suggest_page_categories(d.id), ["CI", "CI", "PL", "PL"])
        with self.assertRaises(ValueError):                   # unassigned page rejected
            self.ws.split(d.id, ["CI", "CI", "PL", None])
        parts = self.ws.split(d.id, ["CI", "CI", "PL", "PL"])
        self.assertEqual(sorted((p.category, p.pages) for p in parts), [("CI", 2), ("PL", 2)])
        self.assertEqual(sum(p.pages for p in parts), 4)
        self.assertTrue(all(p.status == "PDF Ready" for p in parts))
        outs, errs = self.ws.generate("SIN01234567")
        self.assertEqual(errs, [])
        self.assertEqual({o.category: o.pages for o in outs}, {"CI": 2, "PL": 2})
        self.assertIn("Invoice No. 7", self.text_of([o for o in outs if o.category == "CI"][0].path))

    def test_split_interleaved_pages_keep_page_order(self):
        new, _ = run_all(self.ws, "x.pdf", make_pdf([["page one"], ["page two"], ["page three"]]))
        self.ws.split(new[0].id, ["CI", "PL", "CI"])
        outs, _ = self.ws.generate("SIN01234567")
        t = self.text_of([o for o in outs if o.category == "CI"][0].path)
        self.assertLess(t.index("page one"), t.index("page three"))


class TestHousekeeping(Base):
    def test_cleanup_and_retention_purge(self):
        ws2 = Workspace(self.root)
        run_all(ws2, "Invoice.pdf", invoice())
        old = ws2.dir
        os.utime(old, (1, 1))
        import core
        core._last_purge = 0
        removed = Workspace.purge_old(self.root, max_age_s=3600, min_interval_s=0)
        self.assertGreaterEqual(removed, 1)
        self.assertFalse(old.exists())
        self.assertTrue(self.ws.dir.exists())                  # active session untouched
        self.ws.cleanup()
        self.assertFalse(self.ws.dir.exists())

    def test_sessions_are_isolated(self):
        other = Workspace(self.root)
        run_all(self.ws, "Invoice.pdf", invoice())
        self.assertEqual(other.docs, [])
        self.assertNotEqual(other.dir, self.ws.dir)
        other.cleanup()


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
